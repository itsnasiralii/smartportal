# Outage processing rules restored from the supplied NOC reference.
import os, sys, json, io, re, tempfile
from datetime import datetime
from collections import Counter
from zoneinfo import ZoneInfo
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
PKT = ZoneInfo("Asia/Karachi")
outage_history = []
def clean(value):
    """Clean and strip input string safely."""
    return "" if value is None else str(value).strip()

def extract_service_id(label):
    """Extracts short unique service link identifier for concise subjects."""
    value = clean(label)
    if not value:
        return "Service"

    patterns = [
        r"\bDPLC\d+SL\d+\b",
        r"\bDIA\d+SL\d+\b",
        r"\bTurbo\d+SL\d+\b",
        r"\bLNK[A-Z0-9]+\b",
        r"\bMPLS[A-Z0-9-]*\b",
        r"\bPRI[A-Z0-9-]*\b",
        r"\bSIP[A-Z0-9-]*\b",
    ]

    for pattern in patterns:
        match = re.search(pattern, value, flags=re.IGNORECASE)
        if match:
            return match.group(0)

    if "_" in value:
        candidate = value.split("_")[-1].strip()
        if candidate:
            return candidate[:70]

    return value[:70]

def extract_bw_numeric(label):
    match = re.search(r'(\d+(?:\.\d+)?)\s*(?:M|G|K)BPS', str(label).replace('Mpbs', 'Mbps'), re.IGNORECASE)
    if match:
        val = float(match.group(1))
        if 'G' in match.group(0).upper():
            val *= 1000
        elif 'K' in match.group(0).upper():
            val /= 1000
        return val
    return 0

def extract_bw_text(label):
    match = re.search(r'(\d+(?:\.\d+)?)\s*((?:M|G|K)BPS)', str(label).replace('Mpbs', 'Mbps'), re.IGNORECASE)
    if match:
        return f"{match.group(1)} {match.group(2).upper()}"
    return "NA"

def detect_category(label):
    """Streamlined category detection separating BGP DIA from DIA and transport."""
    lbl = str(label).upper()

    # 1. BGP DIA (Prioritized: matches any link containing BGP)
    if "BGP" in lbl:
        return "BGPDIA"

    # 2. Broadband / Access
    if "TURBO" in lbl or "TURBONET" in lbl:
        return "Turbonet"

    # 3. Transport & Dedicated Circuits
    if "MPLS" in lbl or "VPLS" in lbl:
        return "MPLS"
    if "DPLC" in lbl or "EPL" in lbl:
        return "DPLC"
    if "IPLC" in lbl:
        return "IPLC"
    if "M2M" in lbl:
        return "M2M"

    # 4. Standard DIA
    if "DIA" in lbl:
        return "DIA"

    # 5. Voice / Signaling
    if "SIP" in lbl or "PRI" in lbl:
        return "SIP/PRI"

    return "Other"

def extract_client_name(label):
    parts = [p.strip() for p in str(label).strip().split("_") if p.strip()]
    if parts and "ESSCLIENT" in parts[0].upper():
        parts = parts[1:]
    bw_idx = next((i for i, p in enumerate(parts) if re.search(r'\d+(?:\.\d+)?\s*(?:M|G|K)BPS', p.replace('Mpbs', 'Mbps'), re.I)), None)
    c_parts = parts[:bw_idx] if bw_idx is not None else (parts[:-1] if len(parts) > 1 else parts)
    control_words = {"DIA", "MPLS", "DPLC", "M2M", "TURBO", "TURBONET", "BGP", "BGPDIA", "DIABGP", "SIP", "PRI", "SIPPRI", "CENTRAL", "NORTH", "SOUTH"}
    clean_parts = [p for p in c_parts if p.upper() not in control_words and not p.isdigit()]
    return " ".join(clean_parts[:-1] if len(clean_parts) >= 2 else clean_parts).strip()

def extract_unique_id(label):
    parts = str(label).split('_')
    return parts[-1].strip() if len(parts) > 1 else ""

def load_and_clean_data(file_path):
    if file_path.lower().endswith(".csv"):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            lines = f.readlines()
        header_idx = 0
        for i, line in enumerate(lines):
            if "Location Info" in line and "Last Occurred (ST)" in line:
                header_idx = i
                break
        clean_csv_data = "".join(lines[header_idx:])
        return pd.read_csv(io.StringIO(clean_csv_data), sep=",", engine="python")
    else:
        df = pd.read_excel(file_path, header=None)
        header_idx_list = df[df.apply(lambda r: r.astype(str).str.contains("Location Info", case=False).any(), axis=1)].index
        if not header_idx_list.empty:
            header_idx = header_idx_list[0]
            df.columns = df.iloc[header_idx]
            df = df.iloc[header_idx + 1:].reset_index(drop=True)
        return df

def process_outage_file(file_obj):
    """Processes uploaded CSV/Excel alarm dump files and generates BOTH full & simple excel reports."""
    if file_obj is None:
        return "⚠️ Please upload an Excel or CSV file.", None, None, get_history_table(), generate_handover_summary()

    try:
        file_path = os.fspath(file_obj) if isinstance(file_obj, (str, os.PathLike)) else os.fspath(file_obj.name)
        df = load_and_clean_data(file_path)
        df.columns = [str(c).strip() for c in df.columns]

        if "Location Info" not in df.columns or "Last Occurred (ST)" not in df.columns:
            return "❌ Error: 'Location Info' or 'Last Occurred (ST)' column not found.", None, None, get_history_table(), generate_handover_summary()

        processed = []
        for _, row in df.iterrows():
            raw_label = str(row["Location Info"]).strip()
            if pd.isna(raw_label) or raw_label.upper() == "NAN" or "VISIBILITY" in raw_label.upper() or not raw_label.upper().startswith("ESS"):
                continue

            processed.append({
                "Service": raw_label,
                "Bandwidth_Text": extract_bw_text(raw_label),
                "Numeric_BW": extract_bw_numeric(raw_label),
                "Client": extract_client_name(raw_label),
                "Category": detect_category(raw_label),
                "UniqueID": extract_unique_id(raw_label),
                "Last Occurred (ST)": row["Last Occurred (ST)"],
            })

        if not processed:
            return "❌ No valid ESS service records found in file.", None, None, get_history_table(), generate_handover_summary()

        df_unsorted = pd.DataFrame(processed).drop_duplicates(subset=["Service"], keep="first").reset_index(drop=True)
        df_sorted = df_unsorted.sort_values(by="Numeric_BW", ascending=False).reset_index(drop=True)

        df_sorted["Time_Parsed"] = pd.to_datetime(df_sorted["Last Occurred (ST)"], errors="coerce")
        valid_times = df_sorted["Time_Parsed"].dropna()
        outlook_time = valid_times.max().strftime("%Y-%m-%d %I:%M:%S %p") if not valid_times.empty else "Not Found"

        counts = Counter(df_sorted["Category"])
        summary = ", ".join([f"{c}x {s}" for s, c in sorted(counts.items(), key=lambda x: x[1], reverse=True)])

        priority_df = df_sorted[df_sorted["Numeric_BW"] > 250]
        normal_df = df_sorted[df_sorted["Numeric_BW"] <= 250]

        # Generate Console Text Output
        text_lines = []
        text_lines.append("=" * 80)
        text_lines.append(f"{'OUTLOOK READY TEXT':^80}")
        text_lines.append("=" * 80)
        text_lines.append(f"PTN Layer: {len(df_sorted)}x ({summary}) are down")
        text_lines.append("last occured time:")
        text_lines.append(outlook_time)
        text_lines.append("=" * 80 + "\n")

        text_lines.append("Transmission Section\n")
        text_lines.append("Priority Clients")

        for _, row in priority_df.iterrows():
            text_lines.append(row['Service'])

        text_lines.append("\n")

        for _, row in normal_df.iterrows():
            text_lines.append(row['Service'])

        text_lines.append("\nOutage Thread")
        text_lines.append("-" * 50)
        text_lines.append("\n")

        for _, row in df_unsorted.iterrows():
            text_lines.append(row['Service'])

        console_output = "\n".join(text_lines)

        file_time = valid_times.max().strftime("%Y-%m-%d_%I-%M_%p") if not valid_times.empty else datetime.now().strftime("%Y-%m-%d_%I-%M_%p")
        temp_dir = tempfile.mkdtemp(prefix="noc_export_")

        # 1. Full Categorized Report
        out_filepath = os.path.join(temp_dir, f"Outage_Links_{file_time}.xlsx")
        wb = Workbook()
        ws = wb.active
        ws.title = "Categorized Links"

        fill = PatternFill(fill_type="solid", fgColor="D9D9D9")
        font = Font(bold=True)
        border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))

        ws.append(["Part 1: Raw Unsorted Data"])
        ws.cell(row=1, column=1).font = Font(bold=True)
        headers = ["Service", "Bandwidth", "Client", "Category", "UniqueID", "Last Occurred (ST)"]
        ws.append(headers)

        for c in ws[2]:
            c.fill = fill
            c.font = font
            c.border = border

        for _, r in df_unsorted.iterrows():
            ws.append([r["Service"], r["Bandwidth_Text"], r["Client"], r["Category"], r["UniqueID"], r["Last Occurred (ST)"]])

        gap_row = ws.max_row + 3
        ws.cell(row=gap_row, column=1, value="Part 2: Sorted Data for Transmission Rerouting (Descending Bandwidth)")
        ws.cell(row=gap_row, column=1).font = Font(bold=True, italic=True)

        header_row = gap_row + 1
        ws.append(headers)
        for c in ws[header_row]:
            c.fill = fill
            c.font = font
            c.border = border

        for _, r in df_sorted.iterrows():
            bw_val = float(r["Numeric_BW"])
            bw_num = int(bw_val) if bw_val.is_integer() else bw_val
            if bw_num == 0: bw_num = "NA"
            time_fmt = r['Time_Parsed'].strftime("%Y-%m-%d %I:%M:%S %p") if pd.notna(r['Time_Parsed']) else r["Last Occurred (ST)"]
            ws.append([r["Service"], bw_num, r["Client"], r["Category"], r["UniqueID"], time_fmt])

        for col_idx, width in enumerate([75, 15, 35, 15, 25, 25], start=1):
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width

        wb.save(out_filepath)

        # 2. Simple Outage Links Report
        simple_filepath = os.path.join(temp_dir, "Sample Links.xlsx")
        wb_simple = Workbook()
        ws_simple = wb_simple.active
        ws_simple.title = "Outage Links"

        ws_simple.append(["Links"])
        ws_simple.cell(row=1, column=1).font = Font(bold=True)

        for _, r in df_unsorted.iterrows():
            ws_simple.append([r["Service"]])

        ws_simple.column_dimensions['A'].width = 80
        wb_simple.save(simple_filepath)

        # Append to Shift History Log
        now_str = datetime.now(PKT).strftime("%I:%M %p")
        outage_history.append({
            "id": len(outage_history) + 1,
            "processed_time": now_str,
            "file_name": os.path.basename(file_path),
            "total_links": len(df_sorted),
            "summary": summary,
            "priority_count": len(priority_df),
            "occurred_time": outlook_time,
        })


        return console_output, out_filepath, simple_filepath, get_history_table(), generate_handover_summary()

    except Exception as e:
        return f"❌ Processing Error: {str(e)}", None, None, get_history_table(), generate_handover_summary()

def get_history_table():
    rows = []
    for item in reversed(outage_history):
        rows.append([
            f"Outage #{item['id']}",
            item['processed_time'],
            item['total_links'],
            item['priority_count'],
            item['summary'],
            item['occurred_time'],
        ])
    return rows

def generate_handover_summary():
    if not outage_history:
        return "No outages processed in the current shift history."

    total_outages = len(outage_history)
    total_links = sum(item['total_links'] for item in outage_history)
    total_priority = sum(item['priority_count'] for item in outage_history)

    lines = [
        "==================================================",
        "📋 SHIFT HANDOVER SUMMARY / DAILY OUTAGE REPORT",
        "==================================================",
        f"• Total Outages Handled Today: {total_outages}",
        f"• Total Down Links Processed: {total_links}",
        f"• High Priority Links (>250M): {total_priority}",
        "--------------------------------------------------",
        "DETAILS BREAKDOWN:",
    ]

    for item in outage_history:
        lines.append(
            f"• Outage #{item['id']} ({item['processed_time']}): "
            f"{item['total_links']} Links Down ({item['summary']}) | "
            f"Occurred: {item['occurred_time']}"
        )

    lines.append("==================================================")
    return "\n".join(lines)

if __name__ == '__main__':
    import shutil
    output, full, simple, history, handover = process_outage_file(sys.argv[1])
    if not full:
        print(json.dumps({'success': False, 'message': output}))
        sys.exit(0)
    export_dir = os.path.dirname(full)
    for path, name in [(full, 'full.xlsx'), (simple, 'simple.xlsx')]:
        shutil.move(path, os.path.join(sys.argv[2], name))
    shutil.rmtree(export_dir)
    print(json.dumps({'success': True, 'output': output, 'history': outage_history[0]}))
