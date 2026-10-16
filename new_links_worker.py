import os
import sys
import json
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

STANDARD_HEADERS = [
    'SR', 'type', 'nms', 'circuit_name', 'client_name', 'Client_LAT', 'Client_LONG', 
    'Solution (VRF/Non VRF)', 'ISP', 'Master Number', 'Extensions', 'source_City ', 
    'Sub Region', 'Site_latitude', 'Site_longitude', 'Location Status', 'sink_City ', 
    'Comments', 'MSISDN', 'Region', 'Industry Segment', 'Activation_Status', 
    'Account Region', 'band_width', 'GCSS Scope', 'Link Worth', 'Last Mile Medium', 
    'Customer Type', 'Deployment', 'location', 'CMPAK Site ID', 'Buisness Case ID', 
    'last_mile_media', 'Link modification', 'Link modification date', 'Circuit ID', 
    'FLL Link ID', 'FLL Sale ID', 'unique_id', 'Go Ahead Date', 'Site Owner', 
    'Switch Location', 'vlan', 'OTC', 'MRC', 'Link Comissioning Date', 
    'Handed over to CNOC', 'work_order'
]

HIGHLIGHT_COLS = {
    'type', 'nms', 'circuit_name', 'client_name', 'source_City ', 'sink_City ', 
    'band_width', 'location', 'CMPAK Site ID', 'last_mile_media', 'unique_id', 
    'Switch Location', 'vlan', 'work_order'
}

def clean_val(val):
    if val is None:
        return None
    if isinstance(val, str):
        val = val.replace('\xa0', ' ').strip()
        if val == '':
            return None
    return val

def infer_link_type(rec):
    link_type = clean_val(rec.get('type'))
    if link_type:
        return link_type
    
    nms = str(rec.get('nms') or '')
    circuit = str(rec.get('circuit_name') or '')
    
    m = re.search(r'ESS(?:Client|CLient)_([A-Za-z]+)_', nms, re.IGNORECASE)
    if m:
        return m.group(1).upper()
        
    for candidate in ['DIABGP', 'DIA', 'MPLS', 'DPLC', 'PLC']:
        if candidate in circuit.upper() or candidate in nms.upper():
            return candidate
            
    return None

def process_file(input_path, output_dir):
    try:
        wb_raw = openpyxl.load_workbook(input_path, data_only=True)
        ws_raw = wb_raw.active
        raw_rows = list(ws_raw.iter_rows(values_only=True))
    except Exception as e:
        return {'success': False, 'message': f'Failed to parse Excel file: {str(e)}'}
    
    if not raw_rows:
        return {'success': False, 'message': 'The uploaded file is empty.'}
        
    raw_headers = [str(h).strip().lower() if h else f"col_{i}" for i, h in enumerate(raw_rows[0])]
    
    valid_records = []
    skipped_headers = 0
    skipped_blanks = 0
    
    for r in raw_rows:
        if not any(r):
            skipped_blanks += 1
            continue
        first = str(r[0]).strip().lower() if r[0] is not None else ""
        second = str(r[1]).strip().lower() if r[1] is not None else ""
        if first in ("uniquid", "unique_id") or second == "type":
            skipped_headers += 1
            continue
        
        row_dict = {}
        for h, val in zip(raw_headers, r):
            row_dict[h] = clean_val(val)
        valid_records.append(row_dict)
        
    if not valid_records:
        return {'success': False, 'message': 'No valid link data records were found in the uploaded file.'}

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Impacted Links"
    headers = STANDARD_HEADERS
    ws_out.append(headers)
    
    blue_fill = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True)
    for col_idx, h in enumerate(headers, 1):
        cell = ws_out.cell(row=1, column=col_idx)
        cell.font = header_font
        if h in HIGHLIGHT_COLS:
            cell.fill = blue_fill

    preview_rows = []
    for sr_idx, rec in enumerate(valid_records, 1):
        link_type = infer_link_type(rec)
        uid = rec.get('uniquid') or rec.get('unique_id')
        loc = rec.get('igw_location') or rec.get('source_location')
        
        row_data = {
            'SR': sr_idx,
            'type': link_type,
            'nms': rec.get('nms'),
            'circuit_name': rec.get('circuit_name'),
            'client_name': rec.get('client_name'),
            'source_City ': rec.get('source_location'),
            'sink_City ': rec.get('sink_location'),
            'band_width': rec.get('band_width'),
            'location': loc,
            'CMPAK Site ID': rec.get('source_node'),
            'last_mile_media': rec.get('last_mile_media'),
            'unique_id': uid,
            'Switch Location': rec.get('igw_node'),
            'vlan': rec.get('vlan'),
            'Handed over to CNOC': 'No',
            'work_order': rec.get('work_order')
        }
        
        new_row = [row_data.get(h, None) for h in headers]
        ws_out.append(new_row)
        
        preview_rows.append({
            'sr': sr_idx,
            'unique_id': uid or '-',
            'type': link_type or '-',
            'client_name': rec.get('client_name') or '-',
            'circuit_name': rec.get('circuit_name') or '-',
            'band_width': rec.get('band_width') or '-',
            'cmpak_site_id': rec.get('source_node') or '-',
            'switch_location': rec.get('igw_node') or '-',
            'vlan': rec.get('vlan') or '-',
            'media': rec.get('last_mile_media') or '-',
            'work_order': rec.get('work_order') or '-'
        })

    os.makedirs(output_dir, exist_ok=True)
    out_file = os.path.join(output_dir, 'Refine_data.xlsx')
    wb_out.save(out_file)
    
    return {
        'success': True,
        'total': len(valid_records),
        'preview': preview_rows,
        'message': f"Successfully parsed {len(valid_records)} link records (skipped {skipped_headers} repeated headers and {skipped_blanks} blank rows)."
    }

if __name__ == '__main__':
    if len(sys.argv) < 3:
        print(json.dumps({'success': False, 'message': 'Missing input/output arguments.'}))
        sys.exit(1)
        
    input_file = sys.argv[1]
    out_dir = sys.argv[2]
    
    res = process_file(input_file, out_dir)
    print(json.dumps(res))
