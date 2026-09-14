"""Integration coverage against an isolated PHP server and temporary database."""
import http.cookiejar
import json
import os
from pathlib import Path
import re
import socket
import subprocess
import tempfile
import time
import urllib.request
import urllib.error
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent

def main():
    with tempfile.TemporaryDirectory(prefix='noc-test-') as directory:
        with socket.socket() as sock:
            sock.bind(('127.0.0.1', 0))
            port = sock.getsockname()[1]
        env = dict(os.environ, NOC_DATA_DIR=directory, NOC_PASSWORD='test-password', NOC_PYTHON=os.sys.executable)
        server = subprocess.Popen(['php', '-S', f'127.0.0.1:{port}', 'router.php'], cwd=ROOT, env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        base = f'http://127.0.0.1:{port}/'
        browser = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
        def request(path, data=None, headers=None):
            try: return browser.open(urllib.request.Request(base+path, data=data, headers=headers or {}), timeout=45)
            except urllib.error.HTTPError as error: return error
        try:
            for attempt in range(50):
                try:
                    page = request('index.php').read().decode(); break
                except urllib.error.URLError: time.sleep(.1)
            assert 'Sign In' in page
            csrf = re.search(r'name="csrf" value="([a-f0-9]+)"',page)[1]
            from urllib.parse import urlencode
            page = request('index.php',urlencode(dict(action='noc_login',username='nasir',password='test-password',csrf=csrf)).encode()).read().decode()
            assert 'Outage Analyzer' in page
            csrf = re.search(r'window.nocCsrf = "([a-f0-9]+)"', page)[1]
            def api(action, **values):
                result = request('api.php',urlencode(dict(action=action,**values)).encode(), {'X-NOC-CSRF':csrf})
                raw = result.read().decode()
                try: return json.loads(raw)
                except json.JSONDecodeError: raise AssertionError(f'{action}: HTTP {result.status}: {raw}')
            def rows(): return json.load(request('api.php?action=list_complaints'))['data']
            assert request('welcome.sqlite').status == 404
            assert request('api.php',urlencode(dict(action='add_complaints',raw_labels='invalid')).encode()).status == 403
            result=api('add_complaints',raw_labels='ESS_DIA_Example_500Mbps_DIA123SL1\nESS_DIA_Example_500Mbps_DIA123SL1',default_issue='Link is down.',ticket='TKT-1')
            assert result['success'] and 'Skipped 1' in result['message'],result
            complaint=rows()[0]; cid=complaint['id']
            assert not api('generate_closure',complaint_id=cid)['success']
            result=api('generate_opening',complaint_id=cid,issue='Other / Manual',custom_issue='Packet loss on uplink.',time_mode='Manual',manual_time='2026-09-14T10:00')
            assert result['success'] and '10:00 AM' in result['body'],result
            assert not api('generate_closure',complaint_id=cid,time_mode='Manual',manual_time='2026-09-14T09:59')['success']
            result=api('generate_progress',label=complaint['service_label'],status='Team dispatched',ettr='Not Applicable',audience='Team / Vendor',priority='Critical')
            assert 'dispatched' in result['body'] and 'ETTR' not in result['body'] and result['subject'].startswith('CRITICAL'),result
            assert rows()[0]['last_stage']=='Team dispatched'
            result=api('generate_progress',label=complaint['service_label'],status='Testing in progress',ettr='Awaited')
            assert 'ETTR will be shared' in result['body'] and 'is Awaited' not in result['body']
            for scenario in json.loads((ROOT/'noc_options.json').read_text(encoding='utf-8'))['STATS_SCENARIOS']:
                result=api('generate_stats',label=complaint['service_label'],scenario=scenario,audience='Team / Vendor')
                assert result['success'] and result['body'].startswith('Dear Team'),result
                if scenario=='Banking Application Issue': assert 'PCAPdroid' in result['body']
                if scenario=='Turbo Slow Speed': assert result['policy'] and 'policy' not in result['body']
            result=api('generate_customer',label='ESS_Turbo_Test_Turbo123SL1',findings='["Port is up"]',requested_action='Verify last-mile media',priority='Urgent',ticket='ABC')
            assert result['policy'] and 'ABC' in result['subject'] and 'Port is up.' in result['body']
            for fmt in ['A','B','C']:
                result=api('generate_escalation',label=complaint['service_label'],target='Netsat',level='Critical',format=fmt,priority='Critical',ticket='ABC',to='a@example.com; A@example.com',cc='b@example.com')
                assert result['success'] and result['to']=='A@example.com' and '<table' in result['matrix'],result
                if fmt=='A': assert result['subject'].startswith('CRITICAL')
                if fmt=='C': assert '[ABC]' in result['subject']
            assert not api('generate_escalation',label='Test',target='Netsat',to='bad\r\nBcc: bad')['success']
            from email.parser import BytesParser
            from email import policy
            draft = request('api.php',urlencode(dict(action='download_eml',target='Netsat',to='a@example.com',cc='b@example.com',subject='Edited subject — test',body='Edited draft body',csrf=csrf)).encode()).read()
            mail = BytesParser(policy=policy.default).parsebytes(draft)
            assert str(mail['Subject'])=='Edited subject — test' and mail['X-Unsent']=='1'
            assert mail['To']=='a@example.com' and mail['Cc']=='b@example.com'
            assert 'Edited draft body' in mail.get_body(preferencelist=('plain',)).get_content()
            assert '<table' in mail.get_body(preferencelist=('html',)).get_content()
            result=api('generate_closure',complaint_id=cid,found_at='Vendor',root_cause='Fiber Break (Single)',corrective_action='Auto from Root Cause',time_mode='Manual',manual_time='2026-09-14T11:30')
            assert result['success'] and '1 Hour(s) 30 Minute(s)' in result['body'],result
            assert not api('generate_opening',complaint_id=cid,issue='Link is down.')['success']
            assert not api('generate_closure',complaint_id=cid)['success']
            result=api('generate_opening',label='ESS_DIA_Direct_DIA999SL1',issue='Link is down.')
            assert result['success'] and len(rows())==2
            direct=rows()[0]
            result=api('generate_closure',complaint_id=direct['id'],found_at='Customer',root_cause='No Issue Observed',corrective_action='Auto from Root Cause')
            assert result['success'] and 'Service Impact Duration: NA' in result['body']
            dashboard=json.load(request('api.php?action=get_dashboard&filter=Closed%20Only'))
            assert dashboard['data'][0]['duration']=='NA'
            boundary='noc-test-upload'
            csv='Alarm export\nLocation Info,Last Occurred (ST)\nESSClient_BGP_Test_1Gbps_DIA001SL1,2026-09-14 10:00:00\nESSClient_DPLC_Test_100Mbps_DPLC002SL1,2026-09-14 10:05:00\nESSClient_BGP_Test_1Gbps_DIA001SL1,2026-09-14 10:00:00\nESS_VISIBILITY_500Mbps,2026-09-14 10:00:00\n'
            body=(f'--{boundary}\r\nContent-Disposition: form-data; name="action"\r\n\r\nanalyze_outage\r\n--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="sample.csv"\r\nContent-Type: text/csv\r\n\r\n{csv}\r\n--{boundary}--\r\n').encode()
            result=json.load(request('api.php',body,{'Content-Type':'multipart/form-data; boundary='+boundary,'X-NOC-CSRF':csrf}))
            assert result['success'],result
            history=json.load(request('api.php?action=outage_history'))
            assert history['data'][0]['total_links']==2 and history['data'][0]['priority_count']==1,history
            assert 'BGP' in history['handover']
            for kind in ['full','simple']:
                download=request('api.php?action=download_report&token='+result['token']+'&kind='+kind).read()
                path=Path(directory)/(kind+'.xlsx'); path.write_bytes(download)
                wb=load_workbook(path)
                assert wb.active.max_row>=3
                wb.close()
            assert api('reset_outage_history')['success']
            assert not json.load(request('api.php?action=outage_history'))['data']
            assert api('remove_complaint',complaint_id=cid)['success']
            assert len(rows())==1
            print('PASS: authentication, CSRF, complaint lifecycle, manual times, 12 stats scenarios, progress stages, escalation formats, closure duration, outage deduplication, Excel downloads, history and deletion.')
        finally:
            server.terminate(); server.wait(timeout=10)

if __name__=='__main__': main()
